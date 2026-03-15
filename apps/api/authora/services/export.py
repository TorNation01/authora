"""Export service - DOCX, PDF, EPUB, TXT."""

import io
from typing import Any

from fastapi.responses import StreamingResponse


def plain_text_to_tiptap(text: str) -> dict[str, Any]:
    """Convert plain text to TipTap JSON doc."""
    if not text or not text.strip():
        return {"type": "doc", "content": [{"type": "paragraph"}]}
    blocks = []
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        lines = para.split("\n")
        content = []
        for i, line in enumerate(lines):
            if i > 0:
                content.append({"type": "hardBreak"})
            content.append({"type": "text", "text": line})
        blocks.append({"type": "paragraph", "content": content})
    if not blocks:
        return {"type": "doc", "content": [{"type": "paragraph"}]}
    return {"type": "doc", "content": blocks}


def tiptap_to_plain_text(content: dict[str, Any] | None) -> str:
    """Convert TipTap JSON to plain text. Handles None, empty, and edge cases."""
    if content is None or not isinstance(content, dict):
        return ""
    lines = []

    def extract(node: Any) -> None:
        if isinstance(node, dict):
            if "text" in node:
                text = node["text"]
                lines.append(str(text) if text is not None else "")
            elif "content" in node:
                for c in node["content"]:
                    extract(c)
            if node.get("type") == "hardBreak":
                lines.append("\n")
        elif isinstance(node, list):
            for n in node:
                extract(n)

    if "content" in content and isinstance(content["content"], list):
        for node in content["content"]:
            extract(node)
    return "".join(lines).replace("\n\n\n", "\n\n").strip()


def export_txt(chapters: list[dict]) -> bytes:
    """Export to plain text."""
    parts = []
    for ch in chapters:
        title = ch.get("title", "Untitled")
        content = ch.get("content", {})
        text = tiptap_to_plain_text(content)
        parts.append(f"# {title}\n\n{text}")
    return "\n\n".join(parts).encode("utf-8")


def export_docx(chapters: list[dict], book_title: str) -> bytes:
    """Export to DOCX."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(book_title, 0)

    for ch in chapters:
        title = ch.get("title", "Untitled")
        content = ch.get("content", {})
        text = tiptap_to_plain_text(content)

        doc.add_heading(title, level=1)
        for para in text.split("\n\n"):
            if para.strip():
                p = doc.add_paragraph(para)
                p.paragraph_format.space_after = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_pdf(chapters: list[dict], book_title: str) -> bytes:
    """Export to PDF."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Document, Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(book_title, styles["Title"]))
    story.append(Spacer(1, 24))

    for ch in chapters:
        title = ch.get("title", "Untitled")
        content = ch.get("content", {})
        text = tiptap_to_plain_text(content)

        story.append(Paragraph(title.replace("&", "&amp;"), styles["Heading1"]))
        story.append(Spacer(1, 12))
        for para in text.split("\n\n"):
            if para.strip():
                safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, styles["Normal"]))
                story.append(Spacer(1, 6))
        story.append(Spacer(1, 24))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def export_epub(chapters: list[dict], book_title: str, author: str = "Author") -> bytes:
    """Export to EPUB."""
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("authora-export")
    book.set_title(book_title)
    book.set_language("en")
    book.add_author(author)

    epub_chapters = []
    for i, ch in enumerate(chapters):
        title = ch.get("title", "Untitled")
        content = ch.get("content", {})
        text = tiptap_to_plain_text(content)
        html_content = f"<h1>{title}</h1><p>" + text.replace("\n\n", "</p><p>").replace("\n", "<br/>") + "</p>"
        chapter = epub.EpubHtml(title=title, file_name=f"chap_{i}.xhtml", lang="en")
        chapter.content = html_content
        book.add_item(chapter)
        epub_chapters.append(chapter)

    book.toc = epub_chapters
    book.spine = ["nav"] + epub_chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())

    buffer = io.BytesIO()
    epub.write_epub(buffer, book, {})
    buffer.seek(0)
    return buffer.getvalue()


def get_export_filename(
    book_title: str,
    fmt: str,
    *,
    backup_style: bool = False,
    suffix: str | None = None,
) -> str:
    """Generate safe, predictable export filename.
    backup_style: adds -backup-YYYY-MM-DD before extension.
    suffix: e.g. '_outline', '_synopsis' before extension.
    """
    from datetime import date

    safe = "".join(c if c.isalnum() or c in " -_" else "_" for c in (book_title or "manuscript")).strip()[:50]
    if not safe:
        safe = "manuscript"
    ext = fmt.lower()
    if suffix:
        return f"{safe}{suffix}.{ext}"
    if backup_style:
        return f"{safe}-backup-{date.today().isoformat()}.{ext}"
    return f"{safe}.{ext}"
