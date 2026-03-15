"""Extended export service - front/back matter, TOC, structure, publishing prep."""

import io
import zipfile
from dataclasses import dataclass
from typing import Any

from authora.services.export import tiptap_to_plain_text
from authora.services.export_templates import get_template


@dataclass
class ExportParams:
    """Parameters for export."""

    book_title: str
    author_name: str = "Author"
    include_title_page: bool = True
    include_toc: bool = True
    include_acknowledgements: bool = False
    front_matter: str | None = None
    back_matter: str | None = None
    acknowledgements: str | None = None
    dedication: str | None = None
    epigraph: str | None = None
    copyright_notice: str | None = None
    author_bio: str | None = None
    format_style: str = "manuscript"  # manuscript | print | ebook


def _get_node_text(n: dict) -> str:
    """Extract text from a TipTap node."""
    parts = []
    for c in n.get("content", []):
        if isinstance(c, dict):
            if "text" in c:
                parts.append(c["text"])
            elif "content" in c:
                parts.append(_get_node_text(c))
    return "".join(parts)


def _extract_blocks(content: dict[str, Any]) -> list[dict[str, Any]]:
    """Extract structured blocks from TipTap (headings, paragraphs, lists)."""
    blocks = []

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            t = node.get("type")
            if t == "heading":
                level = node.get("attrs", {}).get("level", 1)
                text = _get_node_text(node)
                if text:
                    blocks.append({"type": "heading", "level": level, "text": text})
            elif t == "paragraph":
                text = _get_node_text(node)
                if text or not blocks:
                    blocks.append({"type": "paragraph", "text": text})
            elif t == "blockquote":
                text = _get_node_text(node)
                if text:
                    blocks.append({"type": "blockquote", "text": text})
            elif t == "bulletList":
                for c in node.get("content", []):
                    walk(c)
            elif t == "orderedList":
                for c in node.get("content", []):
                    walk(c)
            elif t == "listItem":
                text = _get_node_text(node)
                if text:
                    blocks.append({"type": "list_item", "text": text})
            elif t == "hardBreak":
                pass
            elif "content" in node:
                for c in node["content"]:
                    walk(c)
        elif isinstance(node, list):
            for n in node:
                walk(n)

    if isinstance(content, dict) and "content" in content:
        for node in content["content"]:
            walk(node)
    return blocks


def _blocks_to_plain_text(blocks: list[dict]) -> str:
    """Convert structured blocks to plain text with heading markers."""
    lines = []
    for b in blocks:
        t = b.get("type", "")
        text = b.get("text", "").strip()
        if not text and t != "paragraph":
            continue
        if t == "heading":
            level = b.get("level", 1)
            prefix = "#" * level + " "
            lines.append(f"{prefix}{text}")
        elif t == "blockquote":
            lines.append(f"> {text}")
        elif t == "list_item":
            lines.append(f"- {text}")
        else:
            lines.append(text)
        if lines:
            lines.append("")
    return "\n".join(lines).replace("\n\n\n", "\n\n")


def tiptap_to_structured_text(content: dict[str, Any]) -> str:
    """Convert TipTap to plain text preserving headings and structure."""
    blocks = _extract_blocks(content)
    return _blocks_to_plain_text(blocks)


def export_outline(chapters: list[dict]) -> bytes:
    """Export chapter outline only (titles + optional summaries)."""
    parts = []
    for i, ch in enumerate(chapters, 1):
        title = ch.get("title", "Untitled")
        content = ch.get("content", {})
        text = tiptap_to_plain_text(content)
        summary = text[:200] + "..." if len(text) > 200 else text
        parts.append(f"{i}. {title}\n   {summary.strip() or '(no content)'}")
    return "\n\n".join(parts).encode("utf-8")


def export_notes_txt(notes: list[dict]) -> bytes:
    """Export notes as plain text."""
    parts = []
    for n in notes:
        title = n.get("title", "Untitled")
        content = n.get("content", "")
        note_type = n.get("note_type", "general")
        source = n.get("source", "")
        tags = n.get("tags", [])
        header = f"# {title} [{note_type}]"
        if source:
            header += f" (Source: {source})"
        if tags:
            header += f" Tags: {', '.join(tags)}"
        parts.append(f"{header}\n\n{content}")
    return "\n\n---\n\n".join(parts).encode("utf-8")


def _build_chapters_data(chapters: list[dict], use_structured: bool = False) -> list[dict]:
    """Build chapters data with optional structured text."""
    out = []
    for ch in chapters:
        title = ch.get("title", "Untitled")
        content = ch.get("content", {})
        if use_structured:
            text = tiptap_to_structured_text(content)
        else:
            text = tiptap_to_plain_text(content)
        out.append({"title": title, "content": content, "text": text})
    return out


def _add_front_matter_docx(doc, params: ExportParams, add_para, add_break):
    """Add structured front matter to DOCX."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    if params.copyright_notice:
        p = add_para(params.copyright_notice)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_break()
    if params.dedication:
        p = add_para(params.dedication)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_break()
    if params.epigraph:
        p = add_para(params.epigraph)
        p.paragraph_format.left_indent = Inches(0.75)
        p.paragraph_format.right_indent = Inches(0.75)
        add_break()
    if params.front_matter:
        for block in params.front_matter.split("\n\n"):
            if block.strip():
                add_para(block.strip())


def export_full_docx(
    chapters: list[dict],
    params: ExportParams,
) -> bytes:
    """Export to DOCX with professional manuscript formatting."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    tpl = get_template(params.format_style)
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(tpl.top_margin_in)
        section.bottom_margin = Inches(tpl.bottom_margin_in)
        section.left_margin = Inches(tpl.left_margin_in)
        section.right_margin = Inches(tpl.right_margin_in)

    style = doc.styles["Normal"]
    font = style.font
    font.name = tpl.font_name
    font.size = Pt(tpl.font_size)
    style.paragraph_format.line_spacing = tpl.line_spacing

    chapters_data = _build_chapters_data(chapters, use_structured=True)

    def add_break():
        doc.add_page_break()

    def add_para(text: str):
        return doc.add_paragraph(text)

    if params.include_title_page:
        doc.add_paragraph()
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.add_run(params.book_title)
        run.bold = True
        run.font.size = Pt(28)
        run.font.name = tpl.font_name
        doc.add_paragraph()
        author_para = doc.add_paragraph(params.author_name)
        author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_para.paragraph_format.space_before = Pt(24)
        for run in author_para.runs:
            run.font.size = Pt(14)
        add_break()

    has_front = bool(params.copyright_notice or params.dedication or params.epigraph or params.front_matter)
    if has_front:
        _add_front_matter_docx(doc, params, add_para, add_break)
        add_break()

    if params.include_toc and len(chapters_data) > 0:
        toc = doc.add_heading("Table of Contents", level=0)
        toc.runs[0].font.size = Pt(16)
        doc.add_paragraph()
        for i, ch in enumerate(chapters_data, 1):
            doc.add_paragraph(f"{i}.\t{ch['title']}", style="Normal")
        add_break()

    for ch in chapters_data:
        doc.add_heading(ch["title"], level=1)
        for block in ch["text"].split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.strip() == tpl.scene_break or block.strip() == "***" or block.strip() == "* * *":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(24)
                p.paragraph_format.space_after = Pt(24)
                run = p.add_run(tpl.scene_break)
                run.font.size = Pt(12)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif block.startswith("#"):
                level = len(block) - len(block.lstrip("#"))
                text = block.lstrip("#").strip()
                doc.add_heading(text, level=min(level, 3))
            elif block.startswith(">"):
                p = doc.add_paragraph(block[1:].strip())
                p.paragraph_format.left_indent = Inches(0.5)
            else:
                p = doc.add_paragraph(block)
                p.paragraph_format.first_line_indent = Inches(tpl.first_line_indent)
                p.paragraph_format.space_after = Pt(12)
        doc.add_paragraph()

    has_back = bool(params.include_acknowledgements and params.acknowledgements) or params.author_bio or params.back_matter
    if has_back:
        add_break()
        if params.include_acknowledgements and params.acknowledgements:
            doc.add_heading("Acknowledgements", level=0)
            for p in params.acknowledgements.split("\n\n"):
                if p.strip():
                    doc.add_paragraph(p.strip())
            add_break()
        if params.author_bio:
            doc.add_heading("About the Author", level=0)
            for p in params.author_bio.split("\n\n"):
                if p.strip():
                    doc.add_paragraph(p.strip())
            add_break()
        if params.back_matter:
            for p in params.back_matter.split("\n\n"):
                if p.strip():
                    doc.add_paragraph(p.strip())

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_full_txt(chapters: list[dict], params: ExportParams) -> bytes:
    """Export to TXT with front/back matter."""
    parts = []
    if params.include_title_page:
        parts.append(params.book_title)
        parts.append("")
        parts.append(params.author_name)
        parts.append("")
        parts.append("=" * 40)
        parts.append("")

    if params.include_toc:
        parts.append("TABLE OF CONTENTS")
        parts.append("")
        for i, ch in enumerate(chapters, 1):
            parts.append(f"  {i}. {ch.get('title', 'Untitled')}")
        parts.append("")
        parts.append("=" * 40)
        parts.append("")

    if params.front_matter:
        parts.append("FRONT MATTER")
        parts.append("")
        parts.append(params.front_matter)
        parts.append("")
        parts.append("=" * 40)
        parts.append("")

    chapters_data = _build_chapters_data(chapters, use_structured=True)
    for ch in chapters_data:
        parts.append(f"# {ch['title']}")
        parts.append("")
        parts.append(ch["text"])
        parts.append("")

    if params.include_acknowledgements and params.acknowledgements:
        parts.append("=" * 40)
        parts.append("ACKNOWLEDGEMENTS")
        parts.append("")
        parts.append(params.acknowledgements)
        parts.append("")

    if params.back_matter:
        parts.append("=" * 40)
        parts.append("BACK MATTER")
        parts.append("")
        parts.append(params.back_matter)

    return "\n".join(parts).encode("utf-8")


def export_full_pdf(chapters: list[dict], params: ExportParams) -> bytes:
    """Export to PDF with front/back matter and professional formatting."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak

    tpl = get_template(params.format_style)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=inch * tpl.top_margin_in,
        bottomMargin=inch * tpl.bottom_margin_in,
        leftMargin=inch * tpl.left_margin_in,
        rightMargin=inch * tpl.right_margin_in,
    )
    styles = getSampleStyleSheet()
    custom = ParagraphStyle(
        name="CustomBody",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=tpl.font_size,
        leading=tpl.font_size * tpl.line_spacing,
        firstLineIndent=inch * tpl.first_line_indent,
        spaceAfter=12,
    )
    story = []

    def safe(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    if params.include_title_page:
        story.append(Paragraph(safe(params.book_title), styles["Title"]))
        story.append(Spacer(1, 24))
        story.append(Paragraph(safe(params.author_name), styles["Normal"]))
        story.append(PageBreak())

    if params.copyright_notice:
        story.append(Paragraph(safe(params.copyright_notice), styles["Normal"]))
        story.append(Spacer(1, 12))
    if params.dedication:
        story.append(Paragraph(f'<para align="center">{safe(params.dedication)}</para>', styles["Normal"]))
        story.append(Spacer(1, 12))
    if params.epigraph:
        story.append(Paragraph(f'<para leftIndent="36" rightIndent="36">{safe(params.epigraph)}</para>', styles["Normal"]))
        story.append(Spacer(1, 12))
    if params.front_matter:
        for p in params.front_matter.split("\n\n"):
            if p.strip():
                story.append(Paragraph(safe(p.strip()), styles["Normal"]))
                story.append(Spacer(1, 6))
        story.append(PageBreak())

    if params.include_toc and chapters:
        story.append(Paragraph("Table of Contents", styles["Heading1"]))
        story.append(Spacer(1, 12))
        for i, ch in enumerate(chapters, 1):
            story.append(Paragraph(f"{i}. {safe(ch.get('title', 'Untitled'))}", styles["Normal"]))
        story.append(PageBreak())

    chapters_data = _build_chapters_data(chapters, use_structured=False)
    for ch in chapters_data:
        story.append(Paragraph(safe(ch["title"]), styles["Heading1"]))
        story.append(Spacer(1, 12))
        for para in ch["text"].split("\n\n"):
            if para.strip():
                story.append(Paragraph(safe(para.strip()), custom))
                story.append(Spacer(1, 6))
        story.append(Spacer(1, 24))

    if params.include_acknowledgements and params.acknowledgements:
        story.append(PageBreak())
        story.append(Paragraph("Acknowledgements", styles["Heading1"]))
        for p in params.acknowledgements.split("\n\n"):
            if p.strip():
                story.append(Paragraph(safe(p.strip()), styles["Normal"]))
                story.append(Spacer(1, 6))
    if params.author_bio:
        story.append(PageBreak())
        story.append(Paragraph("About the Author", styles["Heading1"]))
        for p in params.author_bio.split("\n\n"):
            if p.strip():
                story.append(Paragraph(safe(p.strip()), styles["Normal"]))
                story.append(Spacer(1, 6))
    if params.back_matter:
        for p in params.back_matter.split("\n\n"):
            if p.strip():
                story.append(Paragraph(safe(p.strip()), styles["Normal"]))
                story.append(Spacer(1, 6))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def export_full_epub(chapters: list[dict], params: ExportParams) -> bytes:
    """Export to EPUB with metadata."""
    from authora.services.export import export_epub

    chapters_data = [{"title": ch.get("title", ""), "content": ch.get("content", {})} for ch in chapters]
    return export_epub(chapters_data, params.book_title, params.author_name)


def export_chapter_single(chapter: dict, format: str, book_title: str) -> bytes:
    """Export single chapter."""
    chapters_data = [{"title": chapter.get("title", "Untitled"), "content": chapter.get("content", {})}]
    from authora.services.export import export_txt, export_docx, export_pdf, export_epub

    if format == "txt":
        return export_txt(chapters_data)
    if format == "docx":
        return export_docx(chapters_data, f"{book_title} - {chapter.get('title', 'Chapter')}")
    if format == "pdf":
        return export_pdf(chapters_data, f"{book_title} - {chapter.get('title', 'Chapter')}")
    if format == "epub":
        return export_epub(chapters_data, f"{book_title} - {chapter.get('title', 'Chapter')}")
    raise ValueError(f"Unknown format: {format}")


def export_chapter_summary_sheet(summaries: list[dict[str, str]], book_title: str, author_name: str = "Author") -> bytes:
    """Export chapter summary sheet as DOCX for editor handoff."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    doc.add_heading(f"Chapter Summary Sheet: {book_title}", level=0)
    doc.add_paragraph(f"Author: {author_name}")
    doc.add_paragraph()
    for i, s in enumerate(summaries, 1):
        doc.add_heading(f"{i}. {s.get('title', 'Untitled')}", level=1)
        doc.add_paragraph(s.get("summary", "(no summary)"))
        doc.add_paragraph()
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_synopsis_package(synopsis: str, book_title: str, author_name: str = "Author") -> bytes:
    """Export synopsis as DOCX."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Synopsis", level=0)
    doc.add_paragraph(book_title)
    doc.add_paragraph(author_name)
    doc.add_paragraph()
    for p in synopsis.split("\n\n"):
        if p.strip():
            doc.add_paragraph(p.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_beta_reader_package(
    chapters: list[dict],
    pack: dict[str, Any],
    params: ExportParams,
) -> bytes:
    """Export beta reader package as ZIP: manuscript + synopsis + feedback form."""
    chapters_data = _build_chapters_data(chapters, use_structured=True)
    manuscript = export_full_docx(chapters_data, params)
    synopsis_doc = export_synopsis_package(
        pack.get("synopsis", ""),
        params.book_title,
        params.author_name,
    )
    feedback_lines = ["BETA READER FEEDBACK FORM", "", "Book: " + params.book_title, ""]
    for q in pack.get("feedback_questions", []):
        feedback_lines.append(f"- {q}")
        feedback_lines.append("  Your feedback: _______________________")
        feedback_lines.append("")
    feedback_txt = "\n".join(feedback_lines).encode("utf-8")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        safe_title = "".join(c if c.isalnum() or c in " -_" else "_" for c in params.book_title)[:40]
        zf.writestr(f"{safe_title}_manuscript.docx", manuscript)
        zf.writestr(f"{safe_title}_synopsis.docx", synopsis_doc)
        zf.writestr(f"{safe_title}_feedback_form.txt", feedback_txt)
        summary_txt = "CHAPTER SUMMARIES\n\n"
        for s in pack.get("chapter_summaries", []):
            summary_txt += f"{s.get('title', '')}: {s.get('summary', '')}\n\n"
        zf.writestr(f"{safe_title}_chapter_summaries.txt", summary_txt.encode("utf-8"))
    buffer.seek(0)
    return buffer.getvalue()


def export_ghostwriter_handoff(
    chapters: list[dict],
    pack: dict[str, Any],
    params: ExportParams,
    has_ghostwriter_content: bool = False,
) -> bytes:
    """Export ghostwritten delivery handoff: manuscript + briefs + handoff notes."""
    chapters_data = _build_chapters_data(chapters, use_structured=True)
    manuscript = export_full_docx(chapters_data, params)
    handoff_lines = [
        "GHOSTWRITER DELIVERY HANDOFF",
        "",
        "Book: " + params.book_title,
        "Author: " + params.author_name,
        "",
        "This package contains:",
        "- Full manuscript (DOCX)",
        "- Synopsis",
        "- Chapter summaries",
        "- Author bio draft",
        "- Handoff notes for your editor",
        "",
        "Handoff notes:",
    ]
    for note in pack.get("handoff_notes", []):
        handoff_lines.append(f"- {note}")
    handoff_txt = "\n".join(handoff_lines).encode("utf-8")
    synopsis_doc = export_synopsis_package(
        pack.get("synopsis", ""),
        params.book_title,
        params.author_name,
    )
    summary_doc = export_chapter_summary_sheet(
        pack.get("chapter_summaries", []),
        params.book_title,
        params.author_name,
    )
    bio_lines = ["AUTHOR BIO (DRAFT)", "", pack.get("author_bio", "")]
    bio_txt = "\n".join(bio_lines).encode("utf-8")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        safe_title = "".join(c if c.isalnum() or c in " -_" else "_" for c in params.book_title)[:40]
        zf.writestr(f"{safe_title}_manuscript.docx", manuscript)
        zf.writestr(f"{safe_title}_synopsis.docx", synopsis_doc)
        zf.writestr(f"{safe_title}_chapter_summaries.docx", summary_doc)
        zf.writestr(f"{safe_title}_author_bio.txt", bio_txt)
        zf.writestr(f"{safe_title}_handoff_notes.txt", handoff_txt)
    buffer.seek(0)
    return buffer.getvalue()


def export_formatting_preview_html(
    chapters: list[dict],
    params: ExportParams,
) -> str:
    """Generate HTML preview of export structure for review before download."""
    import html

    def esc(s: str) -> str:
        return html.escape(s) if s else ""

    parts = []
    parts.append("<div class='export-preview'>")
    parts.append(f"<h1>{esc(params.book_title)}</h1>")
    parts.append(f"<p class='author'>{esc(params.author_name)}</p>")
    if params.copyright_notice:
        parts.append(f"<p class='copyright'>{esc(params.copyright_notice)}</p>")
    if params.dedication:
        parts.append(f"<p class='dedication'>{esc(params.dedication)}</p>")
    if params.epigraph:
        parts.append(f"<blockquote class='epigraph'>{esc(params.epigraph)}</blockquote>")
    if params.front_matter:
        fm = "</p><p>".join(esc(p) for p in params.front_matter.split("\n\n") if p.strip())
        parts.append(f"<div class='front-matter'><p>{fm}</p></div>")
    if params.include_toc and chapters:
        parts.append("<h2>Table of Contents</h2><ol>")
        for ch in chapters:
            parts.append(f"<li>{esc(ch.get('title', 'Untitled'))}</li>")
        parts.append("</ol>")
    parts.append("<h2>Chapters</h2>")
    chapters_data = _build_chapters_data(chapters, use_structured=True)
    for ch in chapters_data:
        parts.append(f"<h3>{esc(ch['title'])}</h3>")
        text = ch["text"][:500] + "..." if len(ch["text"]) > 500 else ch["text"]
        parts.append(f"<p>{esc(text).replace(chr(10), '<br/>')}</p>")
    if params.include_acknowledgements and params.acknowledgements:
        parts.append("<h2>Acknowledgements</h2>")
        ack = "</p><p>".join(esc(p) for p in params.acknowledgements.split("\n\n") if p.strip())
        parts.append(f"<p>{ack}</p>")
    if params.author_bio:
        parts.append("<h2>About the Author</h2>")
        bio = "</p><p>".join(esc(p) for p in params.author_bio.split("\n\n") if p.strip())
        parts.append(f"<p>{bio}</p>")
    if params.back_matter:
        parts.append("<h2>Back Matter</h2>")
        bm = "</p><p>".join(esc(p) for p in params.back_matter.split("\n\n") if p.strip())
        parts.append(f"<p>{bm}</p>")
    parts.append("</div>")
    return "".join(parts)


def export_chapters_zip(chapters: list[dict], book_title: str, format: str = "docx") -> bytes:
    """Export each chapter as separate file in a ZIP."""
    from authora.services.export import export_txt, export_docx, export_pdf

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, ch in enumerate(chapters):
            single = [{"title": ch.get("title", f"Chapter {i+1}"), "content": ch.get("content", {})}]
            if format == "txt":
                data = export_txt(single)
            elif format == "docx":
                data = export_docx(single, f"{book_title}_{i+1}")
            elif format == "pdf":
                data = export_pdf(single, f"{book_title}_{i+1}")
            else:
                data = export_txt(single)
            safe_title = "".join(c if c.isalnum() or c in " -_" else "_" for c in ch.get("title", f"ch{i+1}"))[:40]
            zf.writestr(f"{safe_title}.{format}", data)
    buffer.seek(0)
    return buffer.getvalue()
